#!/usr/bin/env python3
"""Generate an app-style Library Contents Claim Report from the final inventory workbook.

Usage: python generate_claim_report.py inventory.xlsx [output.html]
Unverified book titles and prices are deliberately rendered as 'Unverified'.
The verification-status and tentative-title columns are never rendered.
"""
from __future__ import annotations
import sys, html, statistics
from datetime import datetime
from pathlib import Path
import openpyxl

SWEEP_ID = "lib-40b3687c8a"
NONBOOKS = [
    ("I006-001", "bookend", "Metal bookend supporting books on the middle shelf", "Metal"),
    ("I006-002", "bookend", "Metal bookend on the bottom shelf", "Metal"),
    ("I007-003", "furniture", "Red upholstered armchairs", "Fabric/metal"),
    ("I007-004", "furniture", "Office chair with grey fabric", "Fabric/metal"),
    ("I007-005", "furniture", "White metal utility cart/shelf unit", "Metal"),
    ("I008-006", "bookshelf", "Wooden library shelving unit with metal shelf supports", "Wood and metal"),
]
def esc(v): return html.escape("" if v is None else str(v))
def money(v):
    try: return f"${float(v):,.2f}"
    except (TypeError, ValueError): return "Unverified"
def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python generate_claim_report.py inventory.xlsx [output.html]")
    src=Path(sys.argv[1]); out=Path(sys.argv[2]) if len(sys.argv)>2 else src.with_name("Library_Contents_Claim_Report.html")
    wb=openpyxl.load_workbook(src, data_only=True, read_only=True); ws=wb.active
    header_row=None
    for i,row in enumerate(ws.iter_rows(values_only=True),1):
        if row and str(row[0]).strip()=="Book #" and len(row)>7:
            header_row=i; break
    if not header_row: raise SystemExit("Could not find the inventory header row (Book #).")
    headers=[str(x).strip() if x is not None else "" for x in next(ws.iter_rows(min_row=header_row,max_row=header_row,values_only=True))]
    idx={name:i for i,name in enumerate(headers)}
    books=[]
    for row in ws.iter_rows(min_row=header_row+1,values_only=True):
        try: num=int(row[idx["Book #"]])
        except (ValueError,TypeError,KeyError): continue
        status=str(row[idx.get("Verification Status",5)] or "").strip().lower()
        verified=status in {"verified","identified"}
        raw_title=row[idx.get("Book Title",2)]
        title=str(raw_title).strip() if verified and raw_title else "Unverified"
        author=row[idx.get("Author / Contributor",3)] if verified else ""
        section=row[idx.get("Frame / Shelf Section",1)] or "Not recorded"
        category=row[idx.get("Category / Topic",4)] or "Not recorded"
        width=row[idx.get("Est. Spine Width (cm)",6)]
        price=row[idx.get("Est. Amazon Price (USD)",7)]
        books.append({"n":num,"section":section,"title":title,"author":author,"category":category,"width":width if verified else None,"price":price if verified else None,"verified":verified})
    books.sort(key=lambda x:x["n"])
    verified=sum(b['verified'] for b in books); unverified=len(books)-verified
    widths=[float(b['width']) for b in books if b['width'] is not None and isinstance(b['width'],(int,float))]
    prices=[float(b['price']) for b in books if b['price'] is not None and isinstance(b['price'],(int,float)) and b['verified']]
    avg_width=f"{statistics.mean(widths):.2f} cm" if widths else "Not available"
    avg_price=money(statistics.mean(prices)) if prices else "Not available"
    rows=[]
    for b in books:
        title=esc(b['title']); author=esc(b['author']); width=f"{float(b['width']):.2f} cm" if isinstance(b['width'],(int,float)) else ("Unverified" if not b['verified'] else "Not available")
        price=money(b['price']) if b['verified'] else "Unverified"
        rows.append(f"<tr><td>{b['n']:03d}</td><td>{esc(b['section'])}</td><td>{title}</td><td>{author or '—'}</td><td>{esc(b['category'])}</td><td>{width}</td><td>{price}</td></tr>")
    item_rows="".join(f"<tr><td>{esc(i)}</td><td>{esc(cat)}</td><td>{esc(desc)}</td><td>{esc(mat)}</td><td>Model not legible</td><td>Unverified</td></tr>" for i,cat,desc,mat in NONBOOKS)
    review_rows="".join(f"<tr><td>{b['n']:03d}</td><td>Unverified</td><td>Title and price require confirmation from a legible spine or reliable listing.</td></tr>" for b in books if not b['verified'])
    now=datetime.now().strftime("%d %B %Y")
    html_doc=f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Library Contents Claim Report</title><style>
:root{{--navy:#183b5b;--blue:#254d70;--line:#d9dee5;--muted:#667085;--bg:#f4f7fa;--amber:#fff5dc;--amberline:#d99a23}}*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:#17212b;font:14px/1.5 Inter,Arial,sans-serif}}.top{{background:#fff;border-bottom:1px solid var(--line);padding:22px clamp(16px,4vw,48px)}}.brandrow{{display:flex;align-items:center;justify-content:space-between;gap:18px;flex-wrap:wrap}}.brand{{font-size:13px;letter-spacing:.12em;font-weight:800;color:var(--blue);text-transform:uppercase}}.pill{{border:1px solid #c8d9e8;background:#edf5fb;color:#244d6c;padding:5px 10px;border-radius:999px;font-size:12px;font-weight:700}}h1{{font-size:clamp(24px,3vw,34px);line-height:1.2;margin:12px 0 8px;color:var(--navy)}}.meta{{color:var(--muted);font-size:13px}}main{{max-width:1500px;margin:24px auto;padding:0 clamp(14px,3vw,36px)}}.notice{{background:var(--amber);border-left:4px solid var(--amberline);padding:14px 16px;border-radius:4px;margin-bottom:18px}}.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(155px,1fr));gap:12px;margin:16px 0 26px}}.card{{background:#fff;border:1px solid var(--line);border-radius:8px;padding:15px 16px;box-shadow:0 1px 2px #102a4308}}.card .label{{font-size:12px;color:var(--muted);font-weight:700}}.card .value{{font-size:23px;font-weight:750;color:var(--navy);margin-top:4px}}section{{background:#fff;border:1px solid var(--line);border-radius:9px;padding:20px;margin:16px 0}}h2{{font-size:19px;color:var(--navy);margin:0 0 12px}}p{{margin:6px 0 12px}}.tablewrap{{overflow:auto;border:1px solid var(--line);border-radius:6px}}table{{border-collapse:collapse;width:100%;min-width:850px}}th,td{{padding:9px 10px;text-align:left;vertical-align:top;border-bottom:1px solid #e7ebef}}th{{background:#f0f4f8;color:#30475b;font-size:12px;position:sticky;top:0}}tbody tr:nth-child(even){{background:#fbfcfd}}.subtle{{color:var(--muted);font-size:12px}}footer{{color:var(--muted);font-size:12px;padding:18px 0 34px}}@media print{{body{{background:#fff}}.top,section,.card{{box-shadow:none}}.tablewrap{{overflow:visible}}th{{position:static}}}}
</style></head><body><header class="top"><div class="brandrow"><div class="brand">Library Contents Claim Agent</div><span class="pill">Sweep report · Draft</span></div><h1>Library Contents Claim Report</h1><div class="meta">Source ID: <strong>{SWEEP_ID}</strong> &nbsp;·&nbsp; Inventory date: {now} &nbsp;·&nbsp; Source: final Gemini-generated inventory associated with this sweep</div></header><main><div class="notice"><strong>Assessment draft — not a verified insurance valuation.</strong> Inventory includes {len(books)} numbered book records and {len(NONBOOKS)} non-book items. {verified} book titles are identified and {unverified} are unverified. Unverified titles and their prices are shown as <strong>Unverified</strong>. Widths and prices are estimates from the source workbook and are not independently verified replacement values.</div><div class="cards"><div class="card"><div class="label">Books catalogued</div><div class="value">{len(books)}</div></div><div class="card"><div class="label">Unverified titles</div><div class="value">{unverified} <span style="font-size:14px">({unverified/len(books)*100:.1f}%)</span></div></div><div class="card"><div class="label">Non-book items</div><div class="value">{len(NONBOOKS)}</div></div><div class="card"><div class="label">Mean estimated spine width</div><div class="value">{avg_width}</div></div><div class="card"><div class="label">Mean estimated USD price</div><div class="value">{avg_price}</div></div></div>
<section><h2>1. Summary</h2><p>The sweep inventory enumerates {len(books)} book records across the rack/shelf sections recorded in the source workbook. Unverified entries are labelled <strong>Unverified</strong>; prices for those entries are also <strong>Unverified</strong>. Source workbook dimensions and prices remain estimates.</p></section>
<section><h2>2. Room measurements</h2><p>Room dimensions: Not calibrated. The source inventory does not include a verified RoomPlan measurement import. Floor area, gross wall area, and shelf-covered wall area are not calculated. Import a RoomPlan scan or enter tape-measured dimensions in the app to populate this section.</p><p class="subtle">A browser webcam cannot reliably infer metric dimensions without a known scale reference.</p></section>
<section><h2>3. Books ({len(books)})</h2><div class="tablewrap"><table><thead><tr><th>Book ID</th><th>Rack / shelf section</th><th>Title</th><th>Author / contributor</th><th>Category</th><th>Estimated spine width</th><th>Estimated Amazon price (USD)</th></tr></thead><tbody>{''.join(rows)}</tbody></table></div></section>
<section><h2>4. Non-book items ({len(NONBOOKS)})</h2><div class="tablewrap"><table><thead><tr><th>Item ID</th><th>Category</th><th>Description</th><th>Material</th><th>Model</th><th>Price</th></tr></thead><tbody>{item_rows}</tbody></table></div></section>
<section><h2>5. Review queue ({unverified})</h2><p>Confirm the following entries using a clearer evidence frame or a reliable title/edition match before relying on a replacement valuation.</p><div class="tablewrap"><table><thead><tr><th>Book ID</th><th>Title</th><th>Review reason</th></tr></thead><tbody>{review_rows or '<tr><td colspan="3">No unverified titles.</td></tr>'}</tbody></table></div></section>
<section><h2>6. Second-country comparison</h2><p>Not available in the supplied workbook. No second-country prices have been inferred.</p></section><footer>Generated by Library Contents Claim Agent report generator · Sweep ID {SWEEP_ID}. Preserve this draft with the source workbook and evidence frames. Values not supported by source data remain unverified.</footer></main></body></html>'''
    out.write_text(html_doc,encoding='utf-8'); print(f"Wrote {out} ({len(books)} books, {unverified} unverified, {len(NONBOOKS)} non-book items)")
if __name__=='__main__': main()
